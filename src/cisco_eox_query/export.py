"""CSV and XLSX export helpers for :class:`~cisco_eox_query.v5.models.EOXRecord` instances."""

from __future__ import annotations

import csv
import io
import os
from collections.abc import Sequence
from datetime import date
from pathlib import Path
from typing import BinaryIO, TextIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from cisco_eox_query.v5.models import EOXRecord

EXPORT_COLUMNS: tuple[str, ...] = (
    "EOLProductID",
    "ProductIDDescription",
    "ProductBulletinNumber",
    "LinkToProductBulletinURL",
    "EOXExternalAnnouncementDate",
    "EndOfSaleDate",
    "EndOfSWMaintenanceReleases",
    "EndOfSecurityVulSupportDate",
    "EndOfRoutineFailureAnalysisDate",
    "EndOfServiceContractRenewal",
    "LastDateOfSupport",
    "EndOfSvcAttachDate",
    "UpdatedTimeStamp",
    "QueryType",
    "EOXInputValue",
    "MigrationPIDActiveFlag",
    "MigrationInformation",
    "MigrationOption",
    "MigrationProductId",
    "MigrationProductName",
    "MigrationStrategy",
    "MigrationProductInfoURL",
)

EXCEL_ERROR_CODES: frozenset[str] = frozenset(
    ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A")
)

CSV_COLUMNS: tuple[str, ...] = EXPORT_COLUMNS

LIFECYCLE_DATE_COLUMNS: frozenset[str] = frozenset(
    (
        "EndOfSaleDate",
        "EndOfSWMaintenanceReleases",
        "EndOfSecurityVulSupportDate",
        "EndOfRoutineFailureAnalysisDate",
        "EndOfServiceContractRenewal",
        "LastDateOfSupport",
        "EndOfSvcAttachDate",
    )
)

_LIFECYCLE_DATE_INDICES: tuple[int, ...] = tuple(
    index + 1 for index, name in enumerate(EXPORT_COLUMNS) if name in LIFECYCLE_DATE_COLUMNS
)

_BLACK_FILL = PatternFill(fill_type="solid", fgColor="000000")
_WHITE_FONT = Font(color="FFFFFF")
_RED_FILL = PatternFill(fill_type="solid", fgColor="FF0000")
_GREEN_FILL = PatternFill(fill_type="solid", fgColor="00B050")
_YEAR = 365


def export_to_csv(
    records: Sequence[EOXRecord],
    output: str | os.PathLike[str] | TextIO | None = None,
) -> str | None:
    """Serialize ``records`` to CSV and return the text or write it to ``output``.

    The header row is always written, even for an empty ``records`` list, so
    the output is always a valid CSV document. Columns follow the flattened
    ``EXPORT_COLUMNS`` schema (aliased as ``CSV_COLUMNS``):

    - The first 15 columns map directly to ``EOXRecord`` fields; the
      ``QueryType`` column holds the normalized query type
      (e.g. ``product_id``).
    - The seven ``Migration*`` columns map to ``record.migration_details``;
      they are empty cells when ``migration_details`` is ``None``.

    ``date`` values serialize as ISO 8601 via :meth:`datetime.date.isoformat`;
    ``None`` serializes as an empty cell. Rows use ``lineterminator="\\n"``
    for deterministic cross-platform output.

    Args:
        records: The records to serialize.
        output: One of three output modes:

            - ``None``: return the CSV text as a ``str``.
            - A ``str`` or ``os.PathLike``: write UTF-8 CSV to that path
              (using ``newline=""`` as required by the :mod:`csv` docs) and
              return ``None``.
            - A file-like object (``TextIO``): write to it directly and
              return ``None``. The caller's object is not closed.

    Returns:
        The CSV text when ``output`` is ``None``, otherwise ``None``.
    """
    rows = [_flatten(record) for record in records]

    if output is None:
        buffer = io.StringIO()
        _write_csv(buffer, rows)
        return buffer.getvalue()

    if isinstance(output, (str, os.PathLike)):
        with open(Path(output), "w", encoding="utf-8", newline="") as stream:
            _write_csv(stream, rows)
    else:
        _write_csv(output, rows)
    return None


def export_to_xlsx(
    records: Sequence[EOXRecord],
    output: str | os.PathLike[str] | BinaryIO | None = None,
) -> bytes | None:
    """Serialize ``records`` to an XLSX workbook and return bytes or write to ``output``.

    The header row is always written, even for an empty ``records`` list, so
    the output is always a valid workbook. Columns follow the same flattened
    ``EXPORT_COLUMNS`` schema as :func:`export_to_csv`.

    Lifecycle date columns (the seven ``EndOf*``/``LastDateOfSupport``/
    ``EndOfSvcAttachDate`` columns) are highlighted relative to the current
    date: past dates get a solid black fill with white font, today gets a
    solid red fill, dates 365+ days out get a solid green fill, and dates in
    between get a red-yellow-green gradient fill. The metadata date columns
    (``EOXExternalAnnouncementDate``, ``UpdatedTimeStamp``) and all non-date
    columns are never highlighted.

    ``date`` values become real Excel date cells, ``str`` values become text
    cells, and ``None`` becomes an empty cell. String values that start with
    ``=``, ``+``, ``-``, or ``@`` are forced to text so they are never
    interpreted as formulas (CSV-injection-style defense).

    Args:
        records: The records to serialize.
        output: One of three output modes:

            - ``None``: return the workbook bytes.
            - A ``str`` or ``os.PathLike``: write the workbook to that path
              and return ``None``.
            - A binary file-like object (``BinaryIO``): write to it directly
              and return ``None``. The caller's object is not closed.

    Returns:
        The workbook bytes when ``output`` is ``None``, otherwise ``None``.
    """
    rows = [_flatten(record) for record in records]
    workbook = Workbook()
    worksheet = workbook.active

    if output is None:
        buffer = io.BytesIO()
        _populate_worksheet(worksheet, rows)
        workbook.save(buffer)
        return buffer.getvalue()

    if isinstance(output, (str, os.PathLike)):
        _populate_worksheet(worksheet, rows)
        workbook.save(Path(output))
    else:
        _populate_worksheet(worksheet, rows)
        workbook.save(output)
    return None


def _write_csv(stream: TextIO, rows: list[list[str | date | None]]) -> None:
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(CSV_COLUMNS)
    writer.writerows([[_cell(value) for value in row] for row in rows])


def _write_xlsx_rows(worksheet: Worksheet, rows: list[list[str | date | None]]) -> None:
    """Write the header and every data row into ``worksheet``."""
    _write_xlsx_row(worksheet, 1, EXPORT_COLUMNS)
    for row_index, values in enumerate(rows, start=2):
        _write_xlsx_row(worksheet, row_index, values)


def _populate_worksheet(worksheet: Worksheet, rows: list[list[str | date | None]]) -> None:
    """Write all rows into ``worksheet`` and highlight lifecycle date cells."""
    _write_xlsx_rows(worksheet, rows)
    _highlight_lifecycle_dates(worksheet, rows)


def _write_xlsx_row(
    worksheet: Worksheet, row_index: int, values: Sequence[str | date | None]
) -> None:
    """Write one row of typed values into ``worksheet`` at ``row_index``.

    ``None`` values are left unset (empty cells). ``date`` values are
    assigned directly so openpyxl emits real Excel date cells. String values
    are stripped of illegal control characters before assignment so they never
    trigger :class:`~openpyxl.utils.exceptions.IllegalCharacterError`. Values
    that could be interpreted as formulas or as Excel error codes are forced
    to text via :attr:`~openpyxl.cell.cell.Cell.data_type` so they round-trip
    verbatim.
    """
    for column_index, value in enumerate(values, start=1):
        if value is None:
            continue
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell = worksheet.cell(row=row_index, column=column_index, value=value)
        if isinstance(value, str) and (
            (value and value[0] in "=+-@") or value in EXCEL_ERROR_CODES
        ):
            cell.data_type = "s"


def _highlight_lifecycle_dates(
    worksheet: Worksheet, rows: list[list[str | date | None]], today: date | None = None
) -> None:
    """Apply lifecycle-date highlighting to the data rows of ``worksheet``.

    Only the lifecycle date columns in ``_LIFECYCLE_DATE_INDICES`` are styled;
    metadata date columns and non-date columns are left untouched. ``today``
    defaults to :func:`datetime.date.today` and exists so tests can pass a
    deterministic reference date.
    """
    reference = today if today is not None else date.today()
    for row_index, values in enumerate(rows, start=2):
        for column_index in _LIFECYCLE_DATE_INDICES:
            value = values[column_index - 1]
            if not isinstance(value, date):
                continue
            fill, font = _date_highlight(value, reference)
            cell = worksheet.cell(row=row_index, column=column_index)
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font


def _date_highlight(value: date, today: date) -> tuple[PatternFill | None, Font | None]:
    """Return the ``(fill, font)`` pair that highlights ``value`` vs ``today``.

    - ``value < today``: solid black fill, white font.
    - ``value == today``: solid red fill, default font.
    - ``value >= today + 365``: solid green fill, default font.
    - ``0 < days < 365``: red-yellow-green gradient fill, default font.
    """
    delta = (value - today).days
    if delta < 0:
        return _BLACK_FILL, _WHITE_FONT
    if delta == 0:
        return _RED_FILL, None
    if delta >= _YEAR:
        return _GREEN_FILL, None
    return PatternFill(fill_type="solid", fgColor=_rgb_hex(_gradient_color(delta / _YEAR))), None


def _gradient_color(t: float) -> tuple[int, int, int]:
    """Interpolate the red-yellow-green gradient at ``t`` in ``[0, 1)``.

    For ``t < 0.5`` interpolate red ``(255, 0, 0)`` to yellow
    ``(255, 255, 0)``; for ``t >= 0.5`` interpolate yellow to green
    ``(0, 176, 80)``.
    """
    if t < 0.5:
        scale = t / 0.5
        return (255, round(255 * scale), 0)
    scale = (t - 0.5) / 0.5
    return (
        round(255 * (1 - scale)),
        round(255 - (255 - 176) * scale),
        round(80 * scale),
    )


def _rgb_hex(rgb: tuple[int, int, int]) -> str:
    """Format an ``(r, g, b)`` tuple as a ``RRGGBB`` hex string."""
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def _flatten(record: EOXRecord) -> list[str | date | None]:
    details = record.migration_details
    return [
        record.eol_product_id,
        record.product_id_description,
        record.product_bulletin_number,
        record.link_to_product_bulletin_url,
        record.eox_external_announcement_date,
        record.end_of_sale_date,
        record.end_of_sw_maintenance_releases,
        record.end_of_security_vul_support_date,
        record.end_of_routine_failure_analysis_date,
        record.end_of_service_contract_renewal,
        record.last_date_of_support,
        record.end_of_svc_attach_date,
        record.updated_time_stamp,
        record.eox_input_type,
        record.eox_input_value,
        details.pid_active_flag if details else None,
        details.migration_information if details else None,
        details.migration_option if details else None,
        details.migration_product_id if details else None,
        details.migration_product_name if details else None,
        details.migration_strategy if details else None,
        details.migration_product_info_url if details else None,
    ]


def _cell(value: str | date | None) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return value
