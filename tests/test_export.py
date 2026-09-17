from __future__ import annotations

import csv
import io
from datetime import date, datetime

from openpyxl import load_workbook

import cisco_eox_query
from cisco_eox_query import export_to_csv
from cisco_eox_query.export import CSV_COLUMNS, EXPORT_COLUMNS, export_to_xlsx
from cisco_eox_query.v5.models import EOXRecord


def _rows(csv_text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(csv_text)))


def _column_index(name: str) -> int:
    return CSV_COLUMNS.index(name)


def test_header_row_matches_csv_columns(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    text = export_to_csv([record])
    assert text.splitlines()[0] == ",".join(CSV_COLUMNS)


def test_csv_header_uses_query_type():
    record = EOXRecord.model_validate({"EOLProductID": "WIC-1T="})
    header = _rows(export_to_csv([record]))[0]
    assert "QueryType" in header
    assert "EOXInputType" not in header


def test_csv_query_type_column_contains_normalized_value(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    row = _rows(export_to_csv([record]))[1]
    assert row[_column_index("QueryType")] == "product_id"


def test_full_record_row(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    row = _rows(export_to_csv([record]))[1]
    assert row[_column_index("EOLProductID")] == "WIC-1T="
    assert row[_column_index("EndOfSaleDate")] == "2009-12-28"
    assert row[_column_index("MigrationProductId")] == "HWIC-1T="
    assert row[_column_index("MigrationProductInfoURL")] == "https://www.cisco.com"
    assert row[_column_index("QueryType")] == "product_id"


def test_none_values_serialize_as_empty_cells():
    record = EOXRecord.model_validate({"EOLProductID": "WIC-1T="})
    row = _rows(export_to_csv([record]))[1]
    assert row[0] == "WIC-1T="
    assert row[1:] == [""] * (len(CSV_COLUMNS) - 1)


def test_no_migration_details_leaves_migration_cells_empty():
    record = EOXRecord.model_validate({"EOLProductID": "WIC-1T=", "EndOfSaleDate": "2009-12-28"})
    row = _rows(export_to_csv([record]))[1]
    migration_columns = [c for c in CSV_COLUMNS if c.startswith("Migration")]
    for column in migration_columns:
        assert row[_column_index(column)] == ""


def test_special_characters_round_trip():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": 'comma, "double quote" and\nnewline',
        }
    )
    row = _rows(export_to_csv([record]))[1]
    assert row[_column_index("ProductIDDescription")] == 'comma, "double quote" and\nnewline'


def test_formula_trigger_values_serialized_verbatim():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "=1+2",
            "EOXInputValue": '=HYPERLINK("http://evil.example")',
        }
    )
    row = _rows(export_to_csv([record]))[1]
    assert row[_column_index("ProductIDDescription")] == "=1+2"
    assert row[_column_index("EOXInputValue")] == '=HYPERLINK("http://evil.example")'


def test_non_ascii_values_round_trip():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "Café — WAN Interface Card",
        }
    )
    row = _rows(export_to_csv([record]))[1]
    assert row[_column_index("ProductIDDescription")] == "Café — WAN Interface Card"


def test_output_path_absolute_writes_matching_string(eox_record_payload, tmp_path):
    record = EOXRecord.model_validate(eox_record_payload)
    expected = export_to_csv([record])
    target = tmp_path / "out.csv"
    result = export_to_csv([record], output=target)
    assert result is None
    assert target.is_absolute()
    assert target.read_text(encoding="utf-8") == expected


def test_empty_records_only_header():
    text = export_to_csv([])
    assert text == ",".join(CSV_COLUMNS) + "\n"


def test_output_none_returns_str(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    assert isinstance(export_to_csv([record]), str)


def test_output_path_writes_utf8_file(eox_record_payload, tmp_path):
    record = EOXRecord.model_validate(eox_record_payload)
    expected = export_to_csv([record])
    target = tmp_path / "out.csv"
    result = export_to_csv([record], output=target)
    assert result is None
    assert target.read_text(encoding="utf-8") == expected


def test_output_filelike_writes_and_does_not_close(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    stream = io.StringIO()
    result = export_to_csv([record], output=stream)
    assert result is None
    assert not stream.closed
    assert stream.getvalue() == export_to_csv([record])


def test_export_to_csv_in_package_root():
    assert cisco_eox_query.export_to_csv is export_to_csv
    assert "export_to_csv" in cisco_eox_query.__all__


def _load_sheet(records):
    data = export_to_xlsx(records)
    return load_workbook(io.BytesIO(data)).active


def _cell_value(sheet, name):
    return sheet.cell(row=2, column=_column_index(name) + 1).value


def _read_row(sheet, row=1):
    return [
        sheet.cell(row=row, column=column).value for column in range(1, len(EXPORT_COLUMNS) + 1)
    ]


def test_xlsx_header_row_matches_export_columns(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    sheet = _load_sheet([record])
    assert _read_row(sheet) == list(EXPORT_COLUMNS)


def test_xlsx_date_cells_are_real_dates(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    sheet = _load_sheet([record])
    for column, expected in (
        ("EndOfSaleDate", date(2009, 12, 28)),
        ("LastDateOfSupport", date(2014, 12, 27)),
    ):
        cell = sheet.cell(row=2, column=_column_index(column) + 1)
        assert cell.value is not None
        if isinstance(cell.value, datetime):
            assert cell.value.date() == expected
        else:
            assert cell.value == expected
        assert "yy" in cell.number_format.lower()


def test_xlsx_string_cells_are_text(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    sheet = _load_sheet([record])
    cell = sheet.cell(row=2, column=_column_index("EOLProductID") + 1)
    assert cell.value == "WIC-1T="
    assert cell.data_type == "s"


def test_xlsx_none_values_are_empty_cells():
    record = EOXRecord.model_validate({"EOLProductID": "WIC-1T="})
    sheet = _load_sheet([record])
    assert _cell_value(sheet, "EOLProductID") == "WIC-1T="
    for column in EXPORT_COLUMNS[1:]:
        assert sheet.cell(row=2, column=_column_index(column) + 1).value is None


def test_xlsx_migration_details_mapping(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    sheet = _load_sheet([record])
    assert _cell_value(sheet, "MigrationProductId") == "HWIC-1T="
    assert _cell_value(sheet, "MigrationProductInfoURL") == "https://www.cisco.com"


def test_xlsx_query_type_column_normalized(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    sheet = _load_sheet([record])
    assert _cell_value(sheet, "QueryType") == "product_id"


def test_xlsx_formula_trigger_values_are_text_not_formulas():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "+1+2",
            "ProductIDDescription": "=1+2",
            "ProductBulletinNumber": "-1+2",
            "LinkToProductBulletinURL": "@SUM(1,2)",
            "EOXInputValue": '=HYPERLINK("http://evil.example")',
        }
    )
    sheet = _load_sheet([record])
    for column, expected in (
        ("EOLProductID", "+1+2"),
        ("ProductIDDescription", "=1+2"),
        ("ProductBulletinNumber", "-1+2"),
        ("LinkToProductBulletinURL", "@SUM(1,2)"),
        ("EOXInputValue", '=HYPERLINK("http://evil.example")'),
    ):
        cell = sheet.cell(row=2, column=_column_index(column) + 1)
        assert cell.data_type == "s"
        assert cell.value == expected


def test_xlsx_control_characters_are_stripped():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "bad\x00text\x1f",
        }
    )
    sheet = _load_sheet([record])
    assert _cell_value(sheet, "ProductIDDescription") == "badtext"


def test_xlsx_error_code_strings_are_text():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "#REF!",
        }
    )
    sheet = _load_sheet([record])
    cell = sheet.cell(row=2, column=_column_index("ProductIDDescription") + 1)
    assert cell.data_type == "s"
    assert cell.value == "#REF!"


def test_xlsx_empty_string_cell_is_text():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "",
        }
    )
    sheet = _load_sheet([record])
    cell = sheet.cell(row=2, column=_column_index("ProductIDDescription") + 1)
    assert cell.value is None or cell.data_type == "s"


def test_xlsx_empty_records_only_header():
    sheet = _load_sheet([])
    assert sheet.max_row == 1
    assert _read_row(sheet) == list(EXPORT_COLUMNS)


def test_xlsx_output_none_returns_bytes(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    data = export_to_xlsx([record])
    assert isinstance(data, bytes)
    assert data.startswith(b"PK")


def test_xlsx_output_path_writes_valid_workbook(eox_record_payload, tmp_path):
    record = EOXRecord.model_validate(eox_record_payload)
    target = tmp_path / "out.xlsx"
    result = export_to_xlsx([record], output=target)
    assert result is None
    sheet = load_workbook(target).active
    assert _cell_value(sheet, "EOLProductID") == "WIC-1T="


def test_xlsx_output_filelike_writes_and_does_not_close(eox_record_payload):
    record = EOXRecord.model_validate(eox_record_payload)
    stream = io.BytesIO()
    result = export_to_xlsx([record], output=stream)
    assert result is None
    assert not stream.closed
    sheet = load_workbook(io.BytesIO(stream.getvalue())).active
    assert _cell_value(sheet, "EOLProductID") == "WIC-1T="


def test_xlsx_non_ascii_round_trip():
    record = EOXRecord.model_validate(
        {
            "EOLProductID": "WIC-1T=",
            "ProductIDDescription": "Café — WAN Interface Card",
        }
    )
    sheet = _load_sheet([record])
    assert _cell_value(sheet, "ProductIDDescription") == "Café — WAN Interface Card"


def test_export_to_xlsx_in_package_root():
    assert cisco_eox_query.export_to_xlsx is export_to_xlsx
    assert "export_to_xlsx" in cisco_eox_query.__all__


def test_xlsx_columns_match_csv_columns():
    assert EXPORT_COLUMNS == CSV_COLUMNS
